#!/usr/bin/env python3
"""古紙回収：Excel「月別古紙回収量一覧表」からの移行スクリプト。ローカル実行専用。

2026-09-08依頼。全件移行（2016年〜2026年度シート）。

使い方:
    pip install openpyxl
    python3 scripts/import-paper-records-xlsx.py <xlsxパス> [--emit-sql=<出力先.sql>]

  既定は解析と検証だけを行う（dry-run）。--emit-sql を付けると投入用SQLを書き出す。
  他の移行スクリプト（import-*-xml.mjs 等）はNodeで書いているが、.xlsx の読み取りだけは
  依存ライブラリ（openpyxl）が要るためPythonにしている（XML・CSVはNodeの標準機能で読める）。

移行方針:
  - シートは2形式ある
      新形式（2020年〜2026年）: 1行=1回の回収（年/年月/回収予定日/備考/4区分/合計）
      旧形式（2016年〜2019年）: 月×回収日のクロス集計（8行1組で3か月分を横に並べる）
    旧形式は「月」行を起点に、その下の「回収日」行と各区分の行を読む。年をまたぐ判定は
    「月が前の組より小さくなったら翌年」で行う（2016年シートだけ1月始まりで、末尾に
    翌年1〜3月の組が付く構成）
  - 旧形式のシート右端には「平成２８年度合計」のような年度集計ブロックがある。回収日として
    緩く数字を拾うとこれを28日と誤読して年間合計を1回分の実績にしてしまうため、日付セルは
    「N日」または数字のみの形（to_day）に限定する
  - 新形式のシートは年度末より先（翌年度の4月分）まで行があり、次のシートの先頭と日付が
    重なる。重複は「実績のある方」を採り、それでも並ぶ場合はその日付の年度に一致する
    シートの方を採る
  - 備考は2種類ある
      「スキップ」「回収なし」等 → skipped=true（合計・月平均の対象から外す）
      日付そのもの（例: 2020-09-23） → その日に振り替えて回収した記録なので
        「2020/09/23に変更」という読める文言へ直して note に残す
  - 実績も備考も無い週は行を作らない（回収予定日は画面側で自動生成するため。設計書 4-18）

検証:
  各回の4区分の和が、新形式は「合計」列、旧形式は「合計」行と一致することを全件突き合わせる。
  一致しない行があれば異常終了する。
"""
import argparse
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit('openpyxl が必要です: pip install openpyxl')

# 1行=1回の新形式で書かれているシート（それ以外はクロス集計の旧形式）
NEW_FORMAT = {'2020年', '2021年 ', '2022年', '2023年', '2024年', '202５年', '2026年'}

CATS = ('cardboard', 'shredder', 'magazine', 'other')
DB_COLS = {'cardboard': 'cardboard_kg', 'shredder': 'shredder_kg',
           'magazine': 'magazine_kg', 'other': 'other_kg'}
# 旧形式の縦の項目名 → 出力キー
OLD_ROW_KEYS = {'段ボール': 'cardboard', 'シュレッダー': 'shredder',
                'シュレッダ': 'shredder', '雑誌': 'magazine', 'その他': 'other'}
SKIP_WORDS = ('スキップ', '回収なし', '中止')


def sheet_year(name):
    """シート名（全角混じり）から西暦を取り出す"""
    return int(re.search(r'(\d{4})', unicodedata.normalize('NFKC', name)).group(1))


def to_num(v):
    """セルの値を数値へ。空欄・文字列・0以下は None"""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    return round(n, 1) if n > 0 else None


def to_day(v):
    """「５日」「12日」等（全角・半角混在）から日を取り出す。年度集計ブロックは弾く"""
    if v is None:
        return None
    s = unicodedata.normalize('NFKC', str(v)).strip()
    if not re.fullmatch(r'\d{1,2}日?', s):
        return None
    day = int(re.match(r'\d{1,2}', s).group())
    return day if 1 <= day <= 31 else None


def to_month(v):
    """「１０月」等から月を取り出す"""
    if v is None:
        return None
    s = unicodedata.normalize('NFKC', str(v)).strip()
    if not re.fullmatch(r'\d{1,2}月?', s):
        return None
    m = int(re.match(r'\d{1,2}', s).group())
    return m if 1 <= m <= 12 else None


def fiscal_year(iso_date):
    y, m = int(iso_date[:4]), int(iso_date[5:7])
    return y if m >= 4 else y - 1


def parse_new(ws, name):
    """新形式: 4行目以降が1回1行"""
    out = []
    for r in range(4, ws.max_row + 1):
        d = ws.cell(r, 3).value
        if not isinstance(d, (datetime, date)):
            continue
        d = d.date() if isinstance(d, datetime) else d
        note = ws.cell(r, 4).value
        out.append({
            'date': d.isoformat(),
            'cardboard': to_num(ws.cell(r, 5).value),
            'shredder': to_num(ws.cell(r, 6).value),
            'magazine': to_num(ws.cell(r, 7).value),
            'other': to_num(ws.cell(r, 8).value),
            'note': str(note).strip() if note is not None and str(note).strip() else None,
            'excel_total': to_num(ws.cell(r, 9).value) or 0,
            'sheet': name,
        })
    return out


def _old_groups(ws):
    """旧形式の「月」行ごとに (月見出しの列と月, 項目行の位置) を返す"""
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or '').strip() != '月':
            continue
        heads = []
        for c in range(2, ws.max_column + 1):
            m = to_month(ws.cell(r, c).value)
            if m:
                heads.append((c, m))
        if not heads:
            continue
        rows = {}
        for rr in range(r + 1, min(r + 9, ws.max_row + 1)):
            label = str(ws.cell(rr, 1).value or '').strip()
            if label == '月':
                break
            if label == '回収日':
                rows['day'] = rr
            elif label == '合計':
                rows['total'] = rr
            elif label in OLD_ROW_KEYS:
                rows[OLD_ROW_KEYS[label]] = rr
        if 'day' in rows:
            yield heads, rows


def parse_old(ws, name):
    """旧形式: 月×回収日のクロス集計を1回1行へ展開する"""
    out = []
    year = sheet_year(name)
    prev_month = None
    for heads, rows in _old_groups(ws):
        for i, (col, month) in enumerate(heads):
            if prev_month is not None and month < prev_month:
                year += 1  # 12月→1月をまたいだ
            prev_month = month
            end = heads[i + 1][0] if i + 1 < len(heads) else ws.max_column + 1
            for c in range(col, end):
                day = to_day(ws.cell(rows['day'], c).value)
                if not day:
                    continue
                try:
                    d = date(year, month, day)
                except ValueError:
                    print(f'  !! 不正な日付 {name} {year}-{month}-{day}', file=sys.stderr)
                    continue
                rec = {'date': d.isoformat(), 'note': None, 'sheet': name,
                       'excel_total': to_num(ws.cell(rows['total'], c).value) or 0
                       if 'total' in rows else 0}
                for cat in CATS:
                    rec[cat] = to_num(ws.cell(rows[cat], c).value) if cat in rows else None
                out.append(rec)
    return out


def normalize_note(note):
    """備考を (skipped, 表示用の文言) へ直す"""
    if not note:
        return False, None
    s = str(note).strip()
    if any(w in s for w in SKIP_WORDS):
        return True, s
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)  # Excelが日付として持っている＝振替日
    if m:
        return False, f'{m.group(1)}/{m.group(2)}/{m.group(3)}に変更'
    return False, s


def dedupe(recs):
    """シート間で重なった日付を1件に寄せる"""
    by_date = defaultdict(list)
    for r in recs:
        by_date[r['date']].append(r)
    chosen = {}
    for d, rows in by_date.items():
        rows.sort(key=lambda r: (-sum(r[c] or 0 for c in CATS),
                                 0 if sheet_year(r['sheet']) == fiscal_year(d) else 1))
        chosen[d] = rows[0]
    return chosen


def build_sql(records):
    """パイプ区切りの1行1レコードを Postgres 側で分解する形にする（1文で入るように）"""
    def cell(v):
        if v is None:
            return ''
        return str(int(v)) if float(v).is_integer() else str(v)

    lines = []
    for r in records:
        assert '|' not in (r['note'] or ''), r
        lines.append('|'.join([
            r['date'], '1' if r['skipped'] else '0',
            cell(r[DB_COLS['cardboard']]), cell(r[DB_COLS['shredder']]),
            cell(r[DB_COLS['magazine']]), cell(r[DB_COLS['other']]), r['note'] or '',
        ]))
    return (
        'insert into paper_records '
        '(collect_date, skipped, cardboard_kg, shredder_kg, magazine_kg, other_kg, note)\n'
        "select split_part(t,'|',1)::date, split_part(t,'|',2)='1',\n"
        "       nullif(split_part(t,'|',3),'')::numeric, nullif(split_part(t,'|',4),'')::numeric,\n"
        "       nullif(split_part(t,'|',5),'')::numeric, nullif(split_part(t,'|',6),'')::numeric,\n"
        "       nullif(split_part(t,'|',7),'')\n"
        'from unnest(string_to_array($D$\n' + '\n'.join(lines) +
        "\n$D$, E'\\n')) t\nwhere t <> ''\non conflict (collect_date) do nothing;\n"
    )


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('--emit-sql', metavar='PATH', help='投入用SQLの出力先')
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    parsed = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = parse_new(ws, name) if name in NEW_FORMAT else parse_old(ws, name)
        print(f'{name!r}: {len(rows)}行')
        parsed.extend(rows)

    # --- 検証: 各回の4区分の和がExcel側の合計と一致するか ---
    bad = [r for r in parsed if abs(sum(r[c] or 0 for c in CATS) - r['excel_total']) > 0.05]
    if bad:
        for r in bad[:20]:
            print(f'  合計不一致 {r["sheet"]} {r["date"]}: '
                  f'自前={sum(r[c] or 0 for c in CATS)} Excel={r["excel_total"]}', file=sys.stderr)
        sys.exit(f'合計が一致しない行が {len(bad)} 件あります')
    print(f'\n合計チェック: {len(parsed)}行すべて一致')

    # --- 取り込み対象を決める ---
    records = []
    for d, r in sorted(dedupe(parsed).items()):
        skipped, note = normalize_note(r['note'])
        if not any(r[c] for c in CATS) and not note:
            continue  # 実績も備考も無い週は行を作らない
        records.append({'date': d, 'skipped': skipped, 'note': note,
                        **{DB_COLS[c]: r[c] for c in CATS}})

    fy_total, fy_count = defaultdict(float), defaultdict(int)
    for r in records:
        t = 0 if r['skipped'] else sum(r[DB_COLS[c]] or 0 for c in CATS)
        fy_total[fiscal_year(r['date'])] += t
        if t > 0:
            fy_count[fiscal_year(r['date'])] += 1
    print(f'\n取り込み対象 {len(records)}件（{records[0]["date"]} 〜 {records[-1]["date"]}）'
          f' / 中止 {sum(1 for r in records if r["skipped"])}件')
    print('\n年度   合計kg  実績回数')
    for f in sorted(fy_total):
        print(f'  {f}  {fy_total[f]:>7.0f}  {fy_count[f]:>3}')

    if args.emit_sql:
        with open(args.emit_sql, 'w', encoding='utf-8') as fp:
            fp.write(build_sql(records))
        print(f'\nSQLを書き出しました: {args.emit_sql}')
    else:
        print('\n（dry-run。SQLを出すには --emit-sql=<パス> を付けてください）')


if __name__ == '__main__':
    main()
